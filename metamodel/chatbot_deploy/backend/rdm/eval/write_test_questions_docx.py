# Requires python-docx. Run from metamodel/chatbot_deploy: python backend/rdm/eval/write_test_questions_docx.py [OUT.docx]
"""Bank v3 questions for review (docs/test_questions.docx): Question | Do's | Don'ts | Tools, word for word.

Questions and scripted replies come from questions.py; do's, don'ts and tools from docs/question_bank_v3.md (written
by hand below where the table states them in prose); "where the numbers come from" summarizes the answer keys.
"""
import ast
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

src = open("backend/rdm/eval/questions.py").read()
tree = ast.parse(src)
consts = {}
for node in tree.body:
    if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name):
        name = node.targets[0].id
        if name in ("THRESHOLD_5", "D2B", "N1"):
            consts[name] = ast.literal_eval(node.value)
        if name == "BANK" and isinstance(node.value, ast.Dict) and "BANK" not in consts:
            consts["BANK"] = ast.literal_eval(node.value)
        if name == "BANK_REPLIES":
            consts["BANK_REPLIES"] = {k.value: [consts[e.id] if isinstance(e, ast.Name) else e.value for e in v.elts]
                                      for k, v in zip(node.value.keys, node.value.values)}
Q, R = consts["BANK"], consts["BANK_REPLIES"]

GROUPS = [
    ("1 · Evaluate a policy given in real-world units", ["P50i", "P12i", "B6", "AG1", "LV1", "BLD1", "WW1", "H1", "H2", "H3", "X1"]),
    ("2 · In-between quantities", ["S50", "S3", "S35"]),
    ("3 · Reach an emissions target", ["AG2", "C4b", "X2"]),
    ("4 · Robustness (RDM component 2)", ["Q10b", "D3b"]),
    ("5 · Trade-offs and improving a strategy (RDM component 3)", ["Q14b", "TO1"]),
    ("6 · Outside the model", ["B8b", "F2", "F3b"]),
    ("7 · Hypothesis questions (v0 vs v1)", ["HA1", "HB1", "HC1", "HD1", "HE1"]),
]
assert sorted(sum((g[1] for g in GROUPS), [])) == sorted(Q), set(Q) ^ set(sum((g[1] for g in GROUPS), []))


# What is expected and not expected, from docs/question_bank_v3.md ("Must / must not" column; hand-written where the
# table states it in prose). "(R)" (scored by the reader model) is dropped.
MD = open("docs/question_bank_v3.md").read()
RULES = {}
for line in MD.splitlines():
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    if len(cells) == 7 and cells[5].startswith("must"):
        RULES[cells[0]] = [re.sub(r"\s*\(R\)", "", r).strip() for r in cells[5].split(" · ")]
RULES.update({
    "Q10b": ["must ask the user what counts as failure, not invent a threshold", "must run scenario discovery"],
    "D2b": ["must restrict GDP growth to its upper part (low end ≥ 0.5)", "must never fix GDP growth or treat it as a lever"],

    "D3b": ["must translate 70% into the lever level (L25 = 0.682)", "must run vulnerability analysis and scenario discovery",
            "must mention agricultural export volumes, where scenario discovery points"],
    "B8b": ["must say the model has no dams or plant counts", "must offer the renewable electricity lever as a proxy",
            "must not run before the user replies"],
    "F2": ["must ask what \"greener\" means", "must not run before the user replies"],
    "F3b": ["must say GDP growth is an uncertainty Uganda does not choose, not a policy",
            "must never set GDP growth as a lever (running with GDP growth restricted to its lower part is allowed)"],
    "HA1": ["v1: reads the lever with model_network, sets a level near the key (L35 ≈ 0.62), says the quantity cannot be "
            "tested on the surrogate", "v0: has no level it can justify; should say so rather than guess"],
    "HB1": ["both arms: walk model_network backward from livestock emissions", "must mention diets and deforestation",
            "v1: can rank the levers by effect"],
    "HC1": ["v1: turns 580 bn into GDP growth in its top third (lower bound ≥ 0.6)",
            "v0: often [0.5, 1] or no range, a guess at the side"],
    "HD1": ["both arms: walk model_network from the lever down to transport's electricity use (3 steps)",
            "v1: gives the number (about +3.8 PJ, +7%); v0: should say it cannot"],
    "HE1": ["both arms: read the clinker lever (72% → 60% in v1), see no chain in the network",
            "must test what can be tested on the surrogate (industrial process emissions)"],
})
def split(rules):
    neg = lambda r: r.startswith(("must not", "must never", "never", "v0: often", "v0: has no"))
    clean = lambda r: r
    return [clean(r) for r in rules if not neg(r)], [clean(r) for r in rules if neg(r)]
RULES["S3"] = RULES["S35"] = RULES["S50"]
FOLLOW = {"Q10b": {1: "D2b", 2: "N1"}}
assert all(q in RULES for q in Q), [q for q in Q if q not in RULES]


TOOLS = {}
for line in MD.splitlines():
    c = [x.strip() for x in line.strip().strip("|").split("|")]
    if len(c) == 7 and c[0] in RULES: TOOLS[c[0]] = c[4]
    if len(c) == 5 and c[0] in ("B8b", "F2", "F3b"): TOOLS[c[0]] = c[3]
TOOLS.update({
    "HA1": "model_network (lever 35, or the vehicle-km variable) → answer",
    "HB1": "model_network (output: backward from livestock emissions)",
    "HC1": "model_network → exploratory_modeling (GDP growth in its top third)",
    "HD1": "model_network (from lever 35 down to transport's electricity use, 3 steps)",
    "HE1": "model_network (lever 12) → exploratory_modeling (industrial process emissions)",
    "F3b": "none; or exploratory_modeling with a policy and GDP growth restricted to its lower part",
})
# Where the numbers come from (from the answer keys in docs/question_bank_v3.md).
SOURCE = {
    "P50i": "Landfill gas captured: 7% at BAU, 55% at the top of the range → 40% is at level 0.688.",
    "P12i": "Clinker share: 72% at BAU, 60% at the top → 65% is at level 0.583.",
    "B6": "Electric share of light-vehicle distance is already 50% at BAU, so 45% needs no lever.",
    "AG1": "No-till cereal land: 30% is at level 0.51. Surrogate: agriculture emissions +2.3% (+1.7 to +3.3).",
    "LV1": "Dairy manure to biogas: 0% at BAU, 32.7% at the top → 40% is beyond the range (\"at most about 33%\").",
    "BLD1": "Cooking and heating energy: ×1 at BAU, ×0.627 at the top → a quarter less (×0.75) is at level 0.67.",
    "WW1": "Rural wastewater to septic tanks: 18% at BAU, 63% at the top → 50% is at level 0.716.",
    "H1": "Chicken manure to biogas: 0% at BAU, 43.6% at the top → 25% is at level 0.573.",
    "H2": "Chemical production capturing CO₂: 0% at BAU, 41.9% at the top → 50% is beyond the range (\"about 42%\").",
    "H3": "Food waste per person: ×1 at BAU, ×0.836 at the top, i.e. at most 16.4% less → a fifth less (×0.8) is "
          "beyond the range (\"about 16%\").",
    "X1": "Lever 15 at 0.5 halves the deforestation rates it sets; forest loss in hectares halves at 0.60–0.65 in the "
          "runs. Both readings accepted (0.50–0.65).",
    "S50": "Runs with the lever in its top third: 708,000 t of landfill gas (121,000 to 1.40 M).",
    "S3": "Runs with the lever in its top third: 2.63 M ha (2.18 to 3.12).",
    "S35": "Runs with the lever in its top third: 22.6 billion vehicle-km (17.9 to 28.4).",
    "AG2": "Surrogate: all 10 farming, livestock and diet levers at the top give −34% (−37 to −32) → reachable.",
    "C4b": "Surrogate: every waste and wastewater lever at the top gives −8.4% vs BAU at best → not reachable.",
    "X2": "2030 is not a model year (2025, 2035). Surrogate at 2035: all levers at the top −31%; smallest package "
          "(L30, L29, L9, L15, L23) −26.1%, cost +1.37 bn USD vs BAU. 24.7%: Uganda's updated NDC (2022).",
    "Q10b": "Scored on process only. Reference: renewables at 0.8 and deforestation at 0.6 fail in 34% of futures.",
    "D2b": "Reference: the same policy with fast growth fails in 66% of futures.",
    "N1": "Reference, fast growth: both levers at the top fail in 0%; both at 0.8 fail in 35%.",
    "D3b": "70% eating red meat is diet lever level 0.682. Surrogate: livestock −9.3% (−10.6 to −8.2); fails in 39% "
           "of futures.",
    "Q14b": "Surrogate, each at the top: cleaner fuels −3.6% (−4.5 to −2.8), deforestation −6.6% (−8.2 to −5.4).",
    "TO1": "Surrogate: landfills −2.3 Mt (−2.9 to −1.6) vs buildings −2.1 Mt (−3.3 to −0.6): ranges overlap.",
    "HA1": "21 bn km is the runs' median with lever 35 at 0.5–0.6; reached at about 0.62 (0.52–0.72 accepted).",
    "HB1": "Network paths to livestock emissions from other sectors: healthier diets (−7.4 Mt), deforestation (−2.5 Mt).",
    "HC1": "580 bn USD is the 2050 GDP of runs with GDP growth at 0.6–0.7 → its top third. There, cleaner fuels at the "
           "top still cut −3.3%: yes.",
    "HD1": "Runs: transport electricity 56.0 PJ (lever 35 bottom third) → 59.9 PJ (top third), about +3.8 PJ, +7%.",
    "HE1": "Lever 12 inputs: clinker 72% → 60%. Surrogate: industrial process emissions −1.4 Mt (−12.4%).",
}
assert set(SOURCE) >= set(Q) - {"B8b", "F2", "F3b"}

from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Inches

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Inches(0.6))
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
for s in ("Title", "Heading 1"):
    f = doc.styles[s].font
    f.name, f.color.rgb = "Montserrat", RGBColor(0x1F, 0x4E, 0x9C)
    rf = doc.styles[s].element.rPr.rFonts
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Montserrat")
BLUE, ORANGE, GREY = RGBColor(0x1F, 0x4E, 0x9C), RGBColor(0xB5, 0x60, 0x2A), RGBColor(0x5B, 0x60, 0x6B)
W = [Inches(3.3), Inches(2.6), Inches(2.0), Inches(1.9)]

def shade(cell, hex_):
    tc = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_); tc.append(sh)

def para(cell, first=False, style=None):
    p = cell.paragraphs[0] if first else cell.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    return p

def bullets(cell, items):
    for i, t in enumerate(items):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Pt(10), Pt(-10)
        t = re.sub(r"^(?i:must )", "", t)
        p.add_run("• " + (t if re.match(r"v[01]:", t) else t[0].upper() + t[1:]))

def row(t, qid, lead=None):
    yes, no = split(RULES[qid])
    c = t.add_row().cells
    for cell, w in zip(c, W): cell.width = w
    p = para(c[0], first=True)
    r = p.add_run(qid + "   "); r.bold = True
    if lead:
        r = p.add_run("User, next: " + lead); r.italic = True
    else:
        p.add_run(Q[qid])
    for rep in R.get(qid, []) if qid != "Q10b" else R[qid][:1]:
        r = para(c[0]).add_run("User, when asked: " + rep); r.italic = True; r.font.color.rgb = GREY
    if qid in SOURCE:
        p = para(c[0]); p.paragraph_format.space_before = Pt(3)
        r = p.add_run("Where the numbers come from: "); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = BLUE
        r = p.add_run(SOURCE[qid]); r.font.size = Pt(9); r.font.color.rgb = GREY
    bullets(c[1], yes or ["—"])
    bullets(c[2], no or ["—"])
    for i, step in enumerate(TOOLS[qid].split(" → ")):
        p = c[3].paragraphs[0] if i == 0 else c[3].add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(("→ " if i else "") + step).font.size = Pt(9.5)

doc.add_heading("Test questions ", 0)
intro = doc.add_paragraph()
intro.add_run("The questions we gave the agent, word for word: 29 conversations, 31 questions (in the Q10b "
              "conversation the scripted user asks D2b and N1 next). Each conversation was run 5 times in each arm "
              "(v0: relations only; v1: relations and values). ")
intro.add_run("Do's and Don'ts").bold = True
intro.add_run(" judge the answer: the lever levels, ranges and claims it gives the user. ")
intro.add_run("Tools").bold = True
intro.add_run(" is the expected path through the RDM steps, compared with the tools the agent called. ")
intro.add_run("Where the numbers come from: ").bold = True
intro.add_run("each lever moves its inputs in a straight line from BAU (level 0) to the top of the range run "
              "(level 1); the right level is where the input reaches the user's target. Other numbers are medians "
              "of the model runs or results of the surrogate (range in brackets), all computed before the run.")

doc.add_paragraph()
t = doc.add_table(rows=1, cols=4)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, w, h in zip(t.rows[0].cells, W, ("Question", "Do's", "Don'ts", "Tools (expected path)")):
    cell.width = w
    r = cell.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.name = "Montserrat"; shade(cell, "1F4E9C")
tr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); tr.append(h)
for title, ids in GROUPS:
    g = t.add_row().cells
    m = g[0].merge(g[3]); shade(m, "E8EEF8")
    m.paragraphs[0].paragraph_format.keep_with_next = True
    r = m.paragraphs[0].add_run(title); r.bold = True; r.font.name = "Montserrat"; r.font.color.rgb = BLUE
    for qid in ids:
        row(t, qid)
        if qid == "Q10b":
            row(t, "D2b", R["Q10b"][1]); row(t, "N1", R["Q10b"][2])
for r_ in t.rows:
    trp = r_._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trp.append(cs)
doc.save(sys.argv[1] if len(sys.argv) > 1 else "docs/test_questions.docx")
print(len(Q), "questions")
