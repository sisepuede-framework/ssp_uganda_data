"""
Review workbook for the group → output mapping (for Alexa and the experts).

Reads backend/rdm/mapping/maps/<run>.csv and writes docs/lever_x_output_mapping.xlsx:
  README         what each sheet holds and what to review
  Groups         one row per lever / uncertainty: counts per source (formulas) and coverage
  Group_outputs  one row per (group, output) with both sources, rho per year, Review column
  Pairs_LX       every lever × uncertainty pair: outputs each moves and outputs both move
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from backend.rdm.mapping.build_mapping import MAPS_DIR, RHO_MOVES
from backend.rdm.runs import get_run

DOCS = Path(__file__).resolve().parents[3] / "docs"
F = "Arial"
HEAD = PatternFill("solid", fgColor="14213D")
INPUT = PatternFill("solid", fgColor="FFF2CC")
GAP = PatternFill("solid", fgColor="FBEDE8")
LINE = Border(bottom=Side(style="thin", color="CDD3DB"))
MOVES = ["both", "empirical only", "structural only"]


def _header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"


def _finish(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, bold=c.font.bold)
            c.border = LINE


def build(run, out: Path) -> None:
    m = pd.read_csv(MAPS_DIR / f"{run.name}.csv")
    wb = Workbook()

    ws = wb.active
    ws.title = "README"
    text = [
        ("Lever / uncertainty → output mapping", True),
        (f"Run {run.run_id} · {int(m.n_runs.iloc[0]):,} sampled runs · rebuild: "
         "python -m backend.rdm.mapping.build_mapping, then python -m backend.rdm.mapping.review_xlsx", False),
        ("", False),
        ("What it is", True),
        ("For each lever (L) and uncertainty (X): which emissions (table decomposed_emissions) and which variables "
         "(table model_output, the ones that explain the emissions) it moves. Only L and X are on the 0–1 scale; "
         "emissions and variables keep their original units.", False),
        ("Two sources", True),
        ("structural = SISEPUEDE's own transformation → output map (levers only). empirical = the runs: Spearman "
         f"correlation between the group's value and the output, every other input varying; it moves the output "
         f"when the largest |rho| across the anchor years is ≥ {RHO_MOVES}.", False),
        ("status: both · empirical only (moves in the runs but not listed by SISEPUEDE: indirect effect) · "
         "structural only (listed but too weak in the runs) · structural, output constant in runs (never varies).", False),
        ("Sheets", True),
        ("Groups — one row per lever and uncertainty, with counts per status (formulas over Group_outputs).", False),
        ("Group_outputs — every link; rho per year; Review and Comment to fill.", False),
        ("Pairs_LX — every lever × uncertainty pair: outputs the lever moves, outputs the uncertainty moves, "
         "and outputs both move (computed in Python from Group_outputs, statuses both / empirical only / structural only).", False),
        ("What to review (yellow cells)", True),
        ("Group_outputs › Review: keep / drop / unsure. Comment: free text. Example: lever 35 × "
         "vehicle_distance_traveled_trns_road_light_electricity › keep › headline variable.", False),
    ]
    ws.column_dimensions["A"].width = 125
    for i, (t, bold) in enumerate(text, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=F, bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # Group_outputs
    wo = wb.create_sheet("Group_outputs")
    years = [c for c in m.columns if c.startswith("rho_") and c[4:].isdigit()]
    cols = ["group_id", "group_kind", "group_name", "sector", "output_field", "output_label", "output_kind",
            "source_table", "status", "rho_max_abs", "rho_year", "rho_sign", *years]
    _header(wo, cols + ["Review", "Comment"],
            [8, 7, 30, 18, 58, 34, 10, 20, 26, 9, 8, 8, *[8] * len(years), 10, 36])
    dv = DataValidation(type="list", formula1='"keep,drop,unsure"', allow_blank=True)
    wo.add_data_validation(dv)
    for r, row in enumerate(m[cols].itertuples(index=False), 2):
        for c, v in enumerate(row, 1):
            wo.cell(row=r, column=c, value=None if pd.isna(v) else v)
        for c in (len(cols) + 1, len(cols) + 2):
            wo.cell(row=r, column=c).fill = INPUT
        dv.add(f"{get_column_letter(len(cols) + 1)}{r}")
    n = len(m) + 1
    wo.auto_filter.ref = f"A1:{get_column_letter(len(cols) + 2)}{n}"
    _finish(wo)

    # Groups (formulas over Group_outputs)
    wg = wb.create_sheet("Groups", 1)
    gcols = ["Group id", "Kind", "Name", "Sector", "Emissions moved", "Variables moved",
             "both", "empirical only", "structural only", "Coverage"]
    _header(wg, gcols, [8, 6, 32, 22, 11, 11, 9, 11, 11, 22])
    names = (m[["group_id", "group_kind", "group_name", "sector"]].drop_duplicates("group_id")
             .set_index("group_id"))
    from backend.rdm.data.groups import lever_ids, uncertainty_ids
    ids = lever_ids(run) + uncertainty_ids(run)
    rng = lambda col: f"Group_outputs!${col}$2:${col}${n}"
    A, G_, I_ = "A", "G", "I"
    for r, gid in enumerate(ids, 2):
        kind = "L" if gid in lever_ids(run) else "X"
        nm = names.loc[gid] if gid in names.index else None
        wg.cell(row=r, column=1, value=gid)
        wg.cell(row=r, column=2, value=kind)
        wg.cell(row=r, column=3, value=None if nm is None else nm.group_name)
        wg.cell(row=r, column=4, value=None if nm is None else nm.sector)
        moved = "+".join(f'COUNTIFS({rng(A)},A{r},{rng(G_)},"{{k}}",{rng(I_)},"{s}")' for s in MOVES)
        wg.cell(row=r, column=5, value="=" + moved.replace("{k}", "emission"))
        wg.cell(row=r, column=6, value="=" + moved.replace("{k}", "variable"))
        for j, s in enumerate(MOVES, 7):
            wg.cell(row=r, column=j, value=f'=COUNTIFS({rng(A)},A{r},{rng(I_)},"{s}")')
        wg.cell(row=r, column=10, value=f'=IF(AND(E{r}>0,F{r}>0),"emissions and variables",'
                                        f'IF(E{r}>0,"emissions only",IF(F{r}>0,"variables only","moves nothing")))')
    last = len(ids) + 1
    wg.conditional_formatting.add(f"A2:J{last}", FormulaRule(formula=['$J2<>"emissions and variables"'], fill=GAP))
    wg.auto_filter.ref = f"A1:J{last}"
    _finish(wg)

    # Pairs_LX (values computed here)
    wp = wb.create_sheet("Pairs_LX")
    _header(wp, ["Lever id", "Lever", "X id", "Uncertainty", "Outputs the lever moves",
                 "Outputs the uncertainty moves", "Outputs both move"], [8, 32, 6, 34, 14, 16, 12])
    mv = m[m.status.isin(MOVES)]
    sets = mv.groupby("group_id").output_field.agg(set).to_dict()
    r = 2
    for li in lever_ids(run):
        for xj in uncertainty_ids(run):
            sl, sx = sets.get(li, set()), sets.get(xj, set())
            vals = [li, names.group_name.get(li), xj, names.group_name.get(xj), len(sl), len(sx), len(sl & sx)]
            for c, v in enumerate(vals, 1):
                wp.cell(row=r, column=c, value=v)
            r += 1
    wp.auto_filter.ref = f"A1:G{r - 1}"
    _finish(wp)

    wb.calculation.fullCalcOnLoad = True
    wb.save(out)


if __name__ == "__main__":
    run = get_run()
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DOCS / "lever_x_output_mapping.xlsx"
    build(run, out)
    print(f"review workbook → {out}")
