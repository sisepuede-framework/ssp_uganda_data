# Requires python-pptx (not in the project env): pip install python-pptx
# Usage: python backend/rdm/deck/build_deck.py docs/rdm_tool_design_v1.pptx
"""Build the RDM Tool Design deck (draft v1) as a .pptx with python-pptx."""
import math
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = sys.argv[1]

NAVY = RGBColor(0x14, 0x21, 0x3D)
INK = RGBColor(0x1E, 0x2A, 0x3D)
MUTED = RGBColor(0x5A, 0x65, 0x76)
BODY = RGBColor(0x3D, 0x4A, 0x5C)
PAPER = RGBColor(0xF7, 0xF6, 0xF2)
MIST = RGBColor(0xEE, 0xF1, 0xF5)
WHITE = RGBColor(0xFD, 0xFC, 0xF9)
LINE = RGBColor(0xCD, 0xD3, 0xDB)
BLUE = RGBColor(0x2F, 0x66, 0x90)
RUST = RGBColor(0xB4, 0x44, 0x2F)
RUST_DK = RGBColor(0x8A, 0x35, 0x24)
RUST_LT = RGBColor(0xFB, 0xED, 0xE8)
PEACH = RGBColor(0xE8, 0xA8, 0x7C)
COVER_TXT = RGBColor(0xF4, 0xF1, 0xEA)
COVER_SUB = RGBColor(0xC9, 0xD3, 0xE0)

HEAD = "Cambria"
SANS = "Calibri"

W, H = 13.333, 7.5
LM, CW = 0.6, 13.333 - 1.2       # left margin, content width
TOP = 1.45                        # content starts below the title
BOTTOM = 6.85                     # content must end above the footer

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
overflow = []


# ── text helpers ─────────────────────────────────────────────────────────────
def est_lines(text, width_in, pt):
    cpl = max(1, int(width_in * 72 / (0.5 * pt)))
    plain = re.sub(r"\*\*", "", text)
    return sum(max(1, math.ceil(len(p) / cpl)) for p in plain.split("\n"))


def est_h(text, width_in, pt, spacing=1.2, pad=0.12):
    return est_lines(text, width_in, pt) * pt * spacing / 72 + pad


def fill_tf(tf, text, pt, color=INK, font=SANS, bold=False, italic=False, align=PP_ALIGN.LEFT,
            spacing=1.1):
    """Write text into a text frame. '**x**' = bold run, '\n' = new paragraph."""
    tf.word_wrap = True
    for i, para in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for j, seg in enumerate(re.split(r"\*\*", para)):
            if not seg:
                continue
            r = p.add_run()
            r.text = seg
            f = r.font
            f.size, f.name, f.color.rgb = Pt(pt), font, color
            f.bold = bold or (j % 2 == 1)
            f.italic = italic


def text(slide, x, y, w, h, s, pt, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(kw.pop("mx", 0))
    tf.margin_top = tf.margin_bottom = Inches(kw.pop("my", 0.02))
    tf.vertical_anchor = kw.pop("anchor", MSO_ANCHOR.TOP)
    fill_tf(tf, s, pt, **kw)
    return tb


def box(slide, x, y, w, h, fill, line=None, radius=True, line_w=1.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = min(0.12, 0.08 / max(h, 0.1) * 1.5)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def callout(slide, y, s, pt=15, fill=MIST, color=INK, w=CW, x=LM):
    h = est_h(s, w - 0.5, pt, pad=0.3)
    box(slide, x, y, w, h, fill)
    text(slide, x + 0.25, y + 0.15, w - 0.5, h - 0.3, s, pt, color=color, anchor=MSO_ANCHOR.MIDDLE)
    return y + h


# ── slide frame ──────────────────────────────────────────────────────────────
def new_slide(title, footer, bg=PAPER, notes=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    text(s, LM, 0.45, CW, 0.85, title, 32, color=NAVY, font=HEAD, bold=True,
         anchor=MSO_ANCHOR.BOTTOM)
    num = len(prs.slides)
    text(s, LM, 7.0, CW - 0.6, 0.3, f"RDM tool design · draft v1 · {footer}" if footer
         else "RDM tool design · draft v1", 11, color=MUTED)
    text(s, W - LM - 0.5, 7.0, 0.5, 0.3, str(num), 11, color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    s._title = title
    return s


def check(slide, y_end):
    if y_end > BOTTOM + 0.01:
        overflow.append((slide._title, round(y_end, 2)))


# ── tables ───────────────────────────────────────────────────────────────────
def table(slide, y, headers, rows, widths, pt=14, head_pt=14, first_bold=True, x=LM,
          rust_first=None):
    """rust_first: set of row indices whose first cell is drawn in rust."""
    pad_x, pad_y = 0.08, 0.05
    heights = [max(est_h(h, w - 2 * pad_x, head_pt, pad=2 * pad_y + 0.04) for h, w in zip(headers, widths))]
    for r in rows:
        heights.append(max(est_h(c, w - 2 * pad_x, pt, pad=2 * pad_y + 0.04) for c, w in zip(r, widths)))
    total_h = sum(heights)
    shp = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y),
                                 Inches(sum(widths)), Inches(total_h))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for ri, hgt in enumerate(heights):
        tbl.rows[ri].height = Inches(hgt)
    for ri, row in enumerate([headers] + rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(pad_x)
            cell.margin_top = cell.margin_bottom = Inches(pad_y)
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = NAVY
                cell.text_frame.text = ""
                fill_tf(cell.text_frame, val, head_pt, color=COVER_TXT, bold=True)
            else:
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if ri % 2 else MIST
                cell.text_frame.text = ""
                color = RUST if (rust_first and ci == 0 and (ri - 1) in rust_first) else INK
                fill_tf(cell.text_frame, val, pt, color=color, bold=(first_bold and ci == 0))
    return y + total_h


# ═════════════════════════════════════════════════════════════════════════════
# 1 · Cover
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
s._title = "cover"
text(s, LM, 1.7, CW, 0.4, "UGANDA SURROGATE CHATBOT · DMDU 2026", 16, color=PEACH, bold=True)
text(s, LM, 2.25, CW, 2.0, "Building RDM into\nthe agent's tools", 54, color=COVER_TXT, font=HEAD,
     bold=True, spacing=1.0)
text(s, LM, 4.45, 10.5, 1.0, "Design draft v1: the tools step by step, what the experiment varies, and "
     "each question type with and without the new information.", 20, color=COVER_SUB)
text(s, LM, 6.85, CW, 0.3, "Branch rdm-tools · draft v1 · 6 October 2026", 12, color=RGBColor(0x9F, 0xB0, 0xC6))
s.notes_slide.notes_text_frame.text = ("Draft for discussion. The deployed system on the metamodel branch is "
                                       "unchanged; this redesign lives on the rdm-tools branch.")

# 2 · Principles
s = new_slide("Six principles behind the redesign", "")
y = TOP + 0.05
for head, body in [
    ("1 · RDM lives in the tools, not in the prompt", "Every evaluation is an ensemble. No tool returns one number from one future."),
    ("2 · Levers are decided; uncertainties are sampled", "Levers outside the policy sit at 0, a decision not to act. All 13 X are sampled unless the question restricts their range."),
    ("3 · Focus is what is shown, not what moves", "All 13 X move. Vulnerability maps project onto the 2 X that matter most for the policy, or the 2 the user picks."),
    ("4 · One tool, one capability, one data source", "Each tool answers one question. System questions read the raw runs; policy questions run the XGBoost. The agent picks tools; the harness never does."),
    ("5 · Declare before you run", "Each experiment states L, X, M and a success criterion, and the user confirms it in plain language. The declaration records the question category."),
    ("6 · Policy-maker language in, real units out", "The user never needs a 0–1 value. The agent translates objectives into lever levels and uncertainty ranges, and results back into real units."),
]:
    rh = max(est_h(head, 4.6, 17), est_h(body, 6.9, 15)) + 0.08
    ln = s.shapes.add_connector(1, Inches(LM), Inches(y), Inches(LM + CW), Inches(y))
    ln.line.color.rgb = LINE
    text(s, LM, y + 0.1, 4.6, rh, head, 17, color=NAVY, bold=True)
    text(s, LM + 5.0, y + 0.1, 7.1, rh, body, 15, color=BODY)
    y += rh
ln = s.shapes.add_connector(1, Inches(LM), Inches(y), Inches(LM + CW), Inches(y))
ln.line.color.rgb = LINE
check(s, y)

# 3 · Layers
s = new_slide("Two data sources, split by the question", "",
              notes="The rule lives in the tools: each tool has one fixed source, so the agent never decides between "
                    "data and surrogate. Every claim should say where it came from.")
text(s, LM, TOP, CW, 0.4, "**About the system** → real SISEPUEDE runs. **About a specific policy** → the XGBoost.", 18, color=NAVY)
y = table(s, TOP + 0.6, ["Layer", "Answers", "Source", "Tools"], [
    ["Knowledge", "What a transformation does, what drives an output, the lever level that reaches an objective", "Lever → output mapping; the (L, X) object from the raw runs; response map", "get_lever_effects, get_output_drivers, solve_lever_level, explain_failure"],
    ["Experimentation", "How this policy performs, where it fails, how it compares, which levels meet a target", "XGBoost, at run time", "define_experiment, run_experiment, find_vulnerabilities, compare_policies, search_policy"],
    ["Catalog and context", "Names, units, what moves what; country context", "Registry and mapping; inventory", "search_catalog, get_country_context"],
], [2.0, 4.3, 2.9, 2.93])
y += 0.3
for i, (h3, body) in enumerate([("Why not the XGBoost alone", "It has 258 outcome targets and no intermediate variables, so it cannot translate tangible objectives or explain a failure."),
                                ("Why not the data alone", "The 100k runs randomise all 54 levers, so none holds a specific policy. They cannot evaluate one across futures.")]):
    x = LM + i * (CW / 2 + 0.15)
    w = CW / 2 - 0.15
    hh = 0.45 + est_h(body, w - 0.5, 15) + 0.2
    box(s, x, y, w, hh, MIST)
    text(s, x + 0.25, y + 0.15, w - 0.5, 0.4, h3, 17, color=NAVY, bold=True)
    text(s, x + 0.25, y + 0.55, w - 0.5, hh - 0.6, body, 15, color=BODY)
check(s, y + hh)

# 4–5 · Data we have, and the (L, X) object
s = new_slide("The data we have", "",
              notes="Local files: data/ssp/sisepuede_run_2026-05-30t21;35;56.244639. Raw runs: "
                    "s3://sisepuede-data/run_database/<run>/. Every run is linked to its L and X values through "
                    "ATTRIBUTE_PRIMARY (primary_id → design_id, future_id). The mapping review file is "
                    "docs/lever_x_output_mapping.xlsx.")
y = table(s, TOP, ["Source", "What it holds", "Files"], [
    ["LHC design", "The 0–1 value of every group in each of the 100k futures: 54 levers (L) and 13 uncertainties (X)",
     "ATTRIBUTE_LHC_SAMPLES_LEVER_EFFECTS, …_EXOGENOUS_UNCERTAINTIES, ATTRIBUTE_PRIMARY"],
    ["Group definitions", "Which input variables each group sets: 547 in the 54 L groups, 93 in the 13 X groups",
     "VARIABLE_TRAJECTORY_GROUPS_L, …_X"],
    ["Raw SISEPUEDE runs", "Every input and output variable, per run and year (2015–2070), in original units",
     "S3: model_input 145 GB, model_output 111 GB, decomposed_emissions 35 GB, cb 2 GB"],
    ["Lever → output mapping", "Which emissions and variables in model_output each lever moves: 619 emissions and 1,079 variables; 35 of 54 levers complete",
     "transformation_output_variable_map; review file lever_x_output_mapping.xlsx"],
    ["XGBoost surrogate", "Trained on the 67 L and X values → 258 outputs (23 emission categories, costs, benefits, GDP; 6 years)",
     "trained_models/"],
], [2.6, 5.6, 3.93], pt=15)
check(s, y)

s = new_slide("What the agent receives: one object per pair (L_i, X_j)", "rust = the experiment's data",
              notes="Checked on 18 runs of model_output: the electric share of light-duty vehicle-km in 2050 is "
                    "50.0% + 21.9 points × L35 (r = 1.000), whatever the GDP draw. L = 0 reproduces BAU. "
                    "The cells are computed from the raw runs; each holds ~4,000 runs with 5 × 5 bins.")
gw = 6.4
gy = table(s, TOP, ["L_i ↓   X_j →", "Low", "Mid", "High"], [
    ["0 (BAU)", "emission_i\nvar_i", "emission_i\nvar_i", "emission_i\nvar_i"],
    ["0.5", "emission_i\nvar_i", "emission_i\nvar_i", "emission_i\nvar_i"],
    ["1 (max)", "emission_i\nvar_i", "emission_i\nvar_i", "emission_i\nvar_i"],
], [1.6, 1.6, 1.6, 1.6], pt=14)
text(s, LM, gy + 0.1, gw, 0.6, "Each cell: median and p10–p90, in original units, over the runs in that cell. "
     "The other 53 L and 12 X keep varying.", 13, color=MUTED, italic=True)
rx = LM + gw + 0.5
rw = CW - gw - 0.5
pts = ["**All** emissions and variables in model_output that L_i moves; a lever can have many of each",
       "Only L and X are on the 0–1 scale; emissions and variables keep their units (Mt CO₂e, PJ, vehicle-km, ha)",
       "Stored once, as one run-level table; a pair is a query, not a file",
       "**v0 receives the names only:** which emissions and variables move, not how much"]
yy = TOP
for p in pts:
    h = est_h(p, rw - 0.4, 16, pad=0.1)
    box(s, rx, yy + 0.08, 0.14, 0.14, RUST, radius=False)
    text(s, rx + 0.35, yy, rw - 0.4, h, p, 16, color=INK)
    yy += h + 0.18
y = callout(s, max(gy + 0.8, yy) + 0.15,
            "**Example · lever 35 × GDP growth, 2050:** electric share of light-duty vehicle-km = 50% at L = 0 → 72% at "
            "L = 1, whatever GDP; GDP moves the volume (26–38 billion vehicle-km) and transport emissions (7.0–12.5 Mt CO₂e).",
            14, fill=RUST_LT, color=INK)
check(s, y)

# 6 · Categories
s = new_slide("Policy makers ask in their own words; the agent translates", "",
              notes="The system is built for policy makers. Nobody, technical users included, can turn a 0–1 lever "
                    "value into a real magnitude, so translation is part of every flow. A lever-value mode for "
                    "technical users can be a separate tool later.")
y = table(s, TOP, ["Category", "How a policy maker asks", "What the agent must translate", "Main tools"], [
    ["A · Information", "\"What can we do about emissions from deforestation?\"", "Which levers move that emission, and how much", "search_catalog, get_output_drivers"],
    ["B · Evaluate a policy", "\"What if half of private cars were electric by 2050?\"", "\"Half electric\" → lever 35 and a level between 0 and 1", "search_catalog, solve_lever_level, run_experiment"],
    ["C · Reach a target", "\"How do we cut transport emissions 30% by 2050?\"", "\"A 30% cut\" → candidate levers and the levels that reach it", "get_output_drivers, search_policy"],
    ["D · Robustness", "\"Does our renewables plan survive an economic boom?\"", "\"A boom\" → a high range of GDP growth: an uncertainty, not a policy", "define_experiment, find_vulnerabilities, explain_failure"],
    ["E · Compare", "\"Electric cars or stopping deforestation: which is better?\"", "Each option → its levers and levels; which side effects to report", "get_lever_effects, compare_policies"],
    ["F · Outside the model", "\"How many jobs would it create?\" · \"Make Uganda greener\"", "\"Jobs\" → not in the model; \"greener\" → ambiguous", "Decline, or ask for clarification"],
], [2.0, 3.75, 3.85, 2.53], pt=14)
y = callout(s, y + 0.25, "**Results come back the same way:** in real units (vehicle-km, megawatts, hectares, Mt CO₂e), "
            "never as lever values.", 14)
check(s, y)

# 7 · Translation both ways
s = new_slide("Translation runs in both directions", "",
              notes="The XGBoost only understands 0–1 lever values and 0–1 uncertainty positions. The policy maker only "
                    "understands real-world objectives and quantities. The agent's tools sit on both sides of the model.")
steps = [("The policy maker's objective", "\"Half of private cars electric by 2050, even if the economy booms\"", PAPER, LINE),
         ("Translate in", "search_catalog: words → lever 35, electric vehicle-km\nsolve_lever_level: 50% → level\ndefine_experiment: \"boom\" → high GDP range", RUST_LT, RUST),
         ("The model", "XGBoost ensemble over the sampled futures (0–1 levers and uncertainties)", WHITE, BLUE),
         ("Translate out", "Emissions in Mt CO₂e, costs in USD, variables in their own units", RUST_LT, RUST),
         ("The answer", "In the policy maker's terms, with the translation stated so it can be checked", PAPER, LINE)]
bw, gap = 2.17, 0.32
for i, (head, body, fill, line) in enumerate(steps):
    x = LM + i * (bw + gap)
    box(s, x, TOP + 0.5, bw, 3.1, fill, line=line, line_w=1.5)
    text(s, x + 0.15, TOP + 0.65, bw - 0.3, 0.6, head, 16, color=NAVY, bold=True)
    text(s, x + 0.15, TOP + 1.35, bw - 0.3, 2.1, body, 13, color=INK)
    if i < len(steps) - 1:
        arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.05), Inches(TOP + 1.9),
                                 Inches(gap - 0.1), Inches(0.3))
        arr.fill.solid()
        arr.fill.fore_color.rgb = MUTED
        arr.line.fill.background()
y = callout(s, TOP + 3.95, "**What the experiment changes:** how well the rust steps are done. In v0 the agent knows "
            "*which* lever and variables are involved but must guess the level; in v1 it gets the level "
            "and its range across futures.", 15)
check(s, y)

# 8 · Tool set overview
s = new_slide("One tool, one capability: 12 tools in five groups", "",
              notes="Each tool answers one question with one data source and one kind of reply. No question can be "
                    "answered by two tools, and each description says when not to use it. The agent chooses the tools; "
                    "the harness never does.")
y = table(s, TOP, ["Group", "Tools", "Data", "v0 / v1"], [
    ["Process", "declare_plan", "—", "Identical"],
    ["Catalog", "search_catalog", "Registry and lever → output mapping (names, units)", "Identical"],
    ["Knowledge", "get_lever_effects · get_output_drivers · solve_lever_level · explain_failure", "The (L, X) object from the raw runs", "v0: names only\nv1: values"],
    ["Experimentation", "define_experiment · run_experiment · find_vulnerabilities · compare_policies · search_policy", "XGBoost, at run time", "Identical"],
    ["Context", "get_country_context", "Inventory", "Identical"],
], [2.2, 5.1, 3.0, 1.83], pt=15, rust_first={2})
y = callout(s, y + 0.3, "**Two boundaries the descriptions must state.** solve_lever_level solves **one** lever from the data, "
            "approximately; search_policy searches **combinations** with XGBoost ensembles and robustness. "
            "define_experiment frames (no computation); run_experiment computes.", 15)
check(s, y)

# 9 · Harness
s = new_slide("The agent decides; the harness enforces", "the harness never chooses a tool",
              notes="A deterministic orchestrator that fixes the tool sequence per category would remove what we "
                    "measure: whether the information changes the agent's choices. So the agent keeps every decision, "
                    "and the harness only keeps state, checks rules, logs, and asks for a plan.")
cards = [("State", "Stores experiments and ensembles by id, so tools chain without the agent copying data"),
         ("Rules", "Checked in code, not asked in the prompt: no run without confirmation; no X fixed at one value; "
                   "C and D need a success criterion; levers in 0–1"),
         ("Log", "Every turn: plan, tool sequence, arguments, returns, \"why\" claims"),
         ("Plan", "declare_plan before any other tool: the category and the steps the agent intends")]
cw4 = (CW - 0.9) / 4
for i, (h3, body) in enumerate(cards):
    x = LM + i * (cw4 + 0.3)
    box(s, x, TOP, cw4, 2.9, WHITE, line=LINE)
    text(s, x + 0.25, TOP + 0.2, cw4 - 0.5, 0.45, h3, 20, color=NAVY, font=HEAD, bold=True)
    text(s, x + 0.25, TOP + 0.8, cw4 - 0.5, 2.0, body, 14, color=BODY)
y = callout(s, TOP + 3.2, "**Why not an orchestrator that picks the tools:** if code fixed the sequence per category, the "
            "chance of calling a knowledge tool would no longer depend on the information, and the effect we want to "
            "measure would disappear. The plan gives two clean scores instead: classification (plan vs true category) "
            "and adherence (plan vs what the agent did).", 15)
check(s, y)

# 10–13 · Tools in detail
TOOL_HEAD = ["Tool", "Steps the tool follows", "Input data", "Output"]
TOOL_W = [2.35, 4.6, 2.6, 2.58]
s = new_slide("Process, catalog and context tools", "identical in v0 and v1")
y = table(s, TOP, TOOL_HEAD, [
    ["declare_plan", "1 · Record the category (A–F)\n2 · Record the tools it intends to use\n3 · Must come before any other tool",
     "The agent's own reading of the question", "Plan id; logged and compared with what the agent did"],
    ["search_catalog", "1 · Match the user's words to names, descriptions, units\n2 · Search levers, uncertainties, emissions, variables\n3 · Return matches and what moves each output",
     "Feature registry; lever → output mapping; units", "Ids, names and units; for each output, the levers that move it"],
    ["get_country_context", "1 · Match the sector or topic\n2 · Return inventory and context figures",
     "Inventory and country data", "Context figures (GDP, population, inventory)"],
], TOOL_W, pt=14)
check(s, y)

s = new_slide("Knowledge tools: same in v0 and v1, what they return differs", "rust column = v1 only",
              notes="All four read the (L, X) object: for each lever and uncertainty pair, how every mapped emission "
                    "and variable responds. Same name, schema and description in v0 and v1.")
y = table(s, TOP, ["Tool", "Steps the tool follows", "v0 returns", "v1 returns"], [
    ["get_lever_effects(lever, [X], year)", "1 · Read the lever's mapped outputs\n2 · Summarise them across L (and X if given)\n3 · Return in original units",
     "The emissions and variables it moves", "+ their values across L and X: median, p10–p90"],
    ["get_output_drivers(output, year)", "1 · Find the levers and X that move the output\n2 · Rank them by how much",
     "The levers and X that move it, unranked", "+ effect of each, ranked"],
    ["solve_lever_level(output, target, year, [X])", "1 · Find levers that move the output\n2 · Read the output against each lever's level, per X bin\n3 · Return the level that reaches the target",
     "The levers that could reach it; no level", "+ level per X region, its range, reachable or not"],
    ["explain_failure(region, policy)", "1 · Take the failure rule from find_vulnerabilities\n2 · Select runs with similar lever levels\n3 · Compare variables inside vs outside",
     "The variables that could differ", "+ which differ, and by how much"],
], [2.6, 4.4, 2.5, 2.63], pt=13, head_pt=14)
check(s, y)

s = new_slide("Experimentation tools: frame, run, discover", "identical in v0 and v1",
              notes="define_experiment refuses to run without a success criterion for C and D. Its warning 'the chosen X "
                    "explains under 1%' reads the response map; in the experiment it must be removed or v0 "
                    "receives v1's information.")
y = table(s, TOP, TOOL_HEAD, [
    ["define_experiment", "1 · Record the policies (lever levels)\n2 · Turn the conditions into X ranges, rest sampled\n3 · Set metrics and the success criterion\n4 · Restate it in plain language for the user",
     "Feature registry; uncertainties in real units", "Experiment id; plain-language framing to confirm (no 0–1 values)"],
    ["run_experiment", "1 · Draw N futures over the 13 X (fixed seed)\n2 · Predict each policy and BAU on them\n3 · Paired change vs BAU in each future\n4 · Score the criterion; store the ensemble",
     "XGBoost surrogate (258 targets); the experiment", "Median, p10–p90, change vs BAU, share of futures passing; ensemble id"],
    ["find_vulnerabilities", "1 · Label each future pass or fail\n2 · Fit CART over the 13 X\n3 · Rank X; keep the 2 that best split\n4 · Build the X₁ × X₂ pass-rate map",
     "Stored ensemble: X values and outcomes", "Failure rules in real units; X₁ × X₂ map; share left to the other X"],
], TOOL_W, pt=13)
check(s, y)

s = new_slide("Experimentation tools: compare, search", "identical in v0 and v1")
y = table(s, TOP, TOOL_HEAD, [
    ["compare_policies", "1 · Load the ensemble (same futures for all)\n2 · Per policy: robustness, costs, benefits\n3 · Regret in each future vs the best policy\n4 · Flag dominated policies",
     "Stored ensemble with 2+ policies", "Comparison table; regret (max, p90); dominated policies"],
    ["search_policy", "1 · Take the target and candidate levers\n2 · Fix one common ensemble of futures\n3 · Search lever levels (method still open)\n4 · Keep the least ambitious that pass ≥ p%",
     "XGBoost; candidate levers chosen by the agent; constraints", "A few policies that meet the target, with robustness; or the maximum reachable"],
], TOOL_W, pt=14)
y = callout(s, y + 0.35, "**Leak to close:** search_policy must not pick default candidate levers from the data. The agent "
            "names the candidates; otherwise v0 gets the data through the back door.", 15, fill=RUST_LT, color=RUST_DK)
check(s, y)

# 14 · Treatment
s = new_slide("v0 and v1 differ in one thing: what the knowledge tools return", "placebo design", bg=MIST,
              notes="Total effect = P(agent calls the tool) × effect given the call; both are reported (intention to "
                    "treat and per protocol). The agent.py prompt is written from scratch, carries no lever-effect "
                    "information, and is frozen and pre-registered with the tools.")
y = table(s, TOP, ["", "v0 · structure only", "v1 · structure and responses"], [
    ["System prompt (agent.py)", "Identical: role, RDM method, how to use the catalog; no lever effects", "Identical"],
    ["Process, catalog, experimentation and context tools", "Present", "Present"],
    ["Knowledge tools (4): name, schema, description", "Present, identical", "Present, identical"],
    ["Knowledge tools return", "Names only: which emissions, variables or levers are involved", "Their values across L_i and X_j (median, p10–p90), original units"],
], [3.5, 4.3, 4.33], pt=15, rust_first={3})
y = callout(s, y + 0.3, "**Why a placebo, not an absent tool:** with the same tools in both versions, the opportunity to call "
            "them and the text the agent reads are identical, so the effect can be attributed to the information.", 15)
y = callout(s, y + 0.2, "**The prompt never mentions a capability only one version has.** Promising data the v0 agent cannot "
            "reach would measure a broken instruction, not missing information.", 15, fill=RUST_LT, color=RUST_DK)
check(s, y)

# 15 · Activation
s = new_slide("Where the knowledge tools enter the RDM process", "rust = activation point",
              notes="Four activation points: before framing (understand the question), at framing (levels and "
                    "uncertainties), after scenario discovery (explain why), and at search (which levers to try). The "
                    "middle of the process, running and comparing ensembles, is the same in v0 and v1 by design.")
y = table(s, TOP, ["RDM step", "v0 (names only), the agent…", "v1 (with the values), the agent…", "Tool"], [
    ["0 · Read the objective", "Knows which levers and variables are involved", "Also knows how much each one moves the objective", "search_catalog, get_output_drivers"],
    ["1 · Frame (XLRM)", "Picks the lever, guesses the level (0–1)", "Gets the level and its range across futures", "solve_lever_level, get_lever_effects"],
    ["2–4 · Run, discover, compare", "XGBoost ensembles", "The same: no new information here", "—"],
    ["3b · Explain why", "Explains from general knowledge, unbacked", "Cites variables that differ inside vs outside the failure region", "explain_failure"],
    ["5 · Iterate, search", "Searches levers chosen by name", "Searches levers that move the target enough; flags an infeasible target early", "get_output_drivers"],
], [2.5, 3.1, 4.0, 2.53], pt=14, rust_first={0, 1, 3, 4})
y = callout(s, y + 0.3, "**Two leaks to close:** define_experiment's \"X explains under 1%\" warning and search_policy's "
            "default candidates both read the data. Remove them, or v0 receives v1's information.",
            15, fill=RUST_LT, color=RUST_DK)
check(s, y)

# 16–20 · Example tool sequences
MONO = "Courier New"
TRACES = [
    ("A · Information", "What can we do about emissions from deforestation?",
     [("get_output_drivers(\"Deforestation\", 2050)", "The levers and uncertainties that move it, unranked; no sizes", 0),
      ("answer", "Lists levers by name; cannot rank them or say how much", 0)],
     [("get_output_drivers(\"Deforestation\", 2050)", "Reduce DEFORESTATION explains 89% of the variation; full ambition cuts 4.7 Mt. Uncertainties barely matter", 1),
      ("answer", "One lever dominates, with its size; no simulation needed", 0)]),
    ("B · Evaluate a policy", "What if half of private cars were electric by 2050?",
     [("search_catalog(\"private cars electric\")", "Lever 35; electric light-duty vehicle-km", 0),
      ("solve_lever_level(electric share, 50%, 2050)", "Lever 35 can move it; no level. The agent guesses 0.5, which is really 61%", 0),
      ("define_experiment → run_experiment", "Credits the goal with cuts from a bigger policy", 0)],
     [("search_catalog(\"private cars electric\")", "Lever 35; electric light-duty vehicle-km", 0),
      ("solve_lever_level(electric share, 50%, 2050)", "50% at L = 0 (BAU), 72% at L = 1, whatever GDP. The goal is already BAU", 1),
      ("get_lever_effects(35, 2050)", "Full ambition: transport emissions −13%, total only −1%", 1)]),
    ("C · Reach a target", "How do we cut transport emissions 30% by 2050?",
     [("get_output_drivers(\"Transportation\", 2050)", "11 transport levers and GDP growth, unranked; no sizes", 0),
      ("search_policy(candidates picked by name)", "Likely starts with electric vehicles", 0),
      ("search_policy (again)", "More rounds; an unreachable target found late", 0)],
     [("get_output_drivers(\"Transportation\", 2050)", "9.4 Mt. Strongest: non-electric efficiency −4.2, medium-duty fuel −2.4, demand −2.1; light-duty electric −1.2. GDP growth +4.0", 1),
      ("search_policy(candidates=[33, 37, 31])", "A short, informed search", 1),
      ("find_vulnerabilities", "GDP growth expected as the main failure driver", 0)]),
    ("D · Robustness", "Does our renewables plan survive an economic boom?",
     [("define_experiment(plan, GDP growth high)", "Other 12 uncertainties sampled", 0),
      ("run_experiment → find_vulnerabilities", "Pass rate under a boom", 0),
      ("explain_failure(region, plan)", "Names the variables that could differ; \"why\" stays unbacked", 0)],
     [("get_output_drivers(\"Electricity\", 2050)", "The renewables target (−3.3 Mt) and GDP growth (+3.3 Mt) explain the variation equally", 1),
      ("define_experiment → run → find_vulnerabilities", "Same tools; GDP shown as the main axis", 0),
      ("explain_failure(region, plan)", "Which variables differ inside vs outside the region, and by how much", 1)]),
    ("E · Compare", "Electric cars or stopping deforestation: which is better?",
     [("get_lever_effects([35, 15], 2050)", "What each moves, by name; no sizes", 0),
      ("define_experiment → run → compare_policies", "Levels guessed; ranked on emissions only", 0)],
     [("get_lever_effects([35, 15], 2050)", "Full ambition: 35 cuts total emissions 1%, 15 cuts 13%; 15 lowers crop and livestock value", 1),
      ("define_experiment(+ those benefits) → run → compare_policies", "A trade-off, not just a ranking", 0)]),
]
COL_W = (CW - 0.35) / 2
for title, q, without, with_ in TRACES:
    s = new_slide(title, "rust = the step returns values (v1) · every run starts with declare_plan · figures: raw runs and response map, 2050",
                  notes="F questions (outside the model, ambiguous requests) use no knowledge tool: same flow in v0 and v1.")
    text(s, LM, TOP - 0.05, CW, 0.4, f"\"{q}\"", 18, color=BODY, italic=True)
    y0 = TOP + 0.55
    ends = []
    for ci, (head, color, steps) in enumerate([("v0: structure only", MUTED, without),
                                               ("v1: with the (L, X) values", RUST, with_)]):
        x = LM + ci * (COL_W + 0.35)
        text(s, x, y0, COL_W, 0.35, head, 17, color=color, bold=True)
        y = y0 + 0.5
        for call, ret, k in steps:
            tw = COL_W - 0.4
            hh = est_h(call, tw, 13, pad=0) + est_h("→ " + ret, tw, 14, pad=0) + 0.3
            box(s, x, y, COL_W, hh, RUST_LT if k else WHITE, line=RUST if k else LINE, line_w=1.5 if k else 1.0)
            text(s, x + 0.2, y + 0.12, tw, 0.3, call, 13, color=RUST if k else NAVY, font=MONO, bold=True)
            text(s, x + 0.2, y + 0.12 + est_h(call, tw, 13, pad=0) + 0.04, tw, hh - 0.4, "→ " + ret, 14, color=INK)
            y += hh + 0.15
        ends.append(y)
    check(s, max(ends))

# 21 · Expected differences
s = new_slide("What we expect to differ between v0 and v1", "",
              notes="Each effect splits into P(agent calls the tool) × effect when it does; both are reported "
                    "(intention to treat and per protocol).")
y = table(s, TOP, ["Dimension", "v0: structure only", "v1: with the values", "Scored by"], [
    ["Lever choice", "By name: intuitive but weak levers", "Levers that actually move the objective", "Overlap with the experts and the data's top drivers"],
    ["Lever level", "Guessed", "Solved, with a range", "Distance to the level the runs say reaches the objective"],
    ["Hitting the objective", "May miss it", "Meets it in the stated share of futures", "Gap between objective and ensemble result"],
    ["Infeasible targets", "Found late or missed", "Reported early, with the maximum reachable", "Correct feasible / infeasible call"],
    ["Efficiency", "Simulations for lookups; more searches", "Lookups from the data; fewer runs", "Tool calls and ensembles per question"],
    ["Side effects", "Emissions only", "Costs and benefits the policy moves", "Side effects named vs the data"],
    ["Explanations", "\"Why\" from general knowledge", "\"Why\" backed by data", "Share of claims backed by explain_failure"],
], [2.3, 3.2, 3.35, 3.28], pt=13, head_pt=13)
y = callout(s, y + 0.2, "**No difference expected (checks):** the plan's category, out-of-model and ambiguous questions, "
            "ensembles instead of single runs, confirmation before running.", 13)
check(s, y)

# Open items
s = new_slide("Open items", "")
items = ["Which value an X takes when it is not varied: ask the team's expert.",
         "Complete the lever → output mapping (19 levers incomplete), map X → outputs, add variable units.",
         "Write the agent.py prompt from scratch (identical in v0 and v1), freeze it with the tools, pre-register.",
         "Build the harness: state, rule checks, logging, declare_plan.",
         "Real units for the 13 uncertainties, so \"a boom\" maps to a range the user can check.",
         "Close the two leaks: define_experiment's weak-X warning and search_policy's default candidates.",
         "The search method behind search_policy.",
         "Recompute the response map from real runs and compare it with the XGBoost version, as a fidelity check."]
y = TOP + 0.05
for i, it in enumerate(items):
    h = est_h(it, CW - 0.7, 17, pad=0.1)
    box(s, LM, y + 0.06, 0.38, 0.38, NAVY)
    text(s, LM, y + 0.06, 0.38, 0.38, str(i + 1), 13, color=COVER_TXT, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, LM + 0.6, y + 0.06, CW - 0.7, h, it, 17, color=INK)
    y += max(h, 0.45) + 0.12
check(s, y)

prs.save(OUT)
print("slides:", len(prs.slides))
print("overflow:", overflow or "none")
