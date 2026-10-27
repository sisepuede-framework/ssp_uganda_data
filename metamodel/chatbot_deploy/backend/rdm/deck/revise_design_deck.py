"""Apply the storyline review (docs/dmdu_talk_storyline_review.docx) to the Claude Design deck.

    python revise_design_deck.py ~/Downloads/DMDU_2026_Talk.pptx ~/Downloads/DMDU_2026_Talk_v2.pptx
"""
import copy
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
S = list(prs.slides)                       # original order, 0-based

BLUE, INK, MUTED, MUTED2 = "1F4E9C", "16181D", "6B707B", "5B606B"
GREY, RULE, ORANGE, LIGHT, PALE, PEACH = "A3A9B4", "E3E5E9", "B5602A", "E8EEF8", "F4F5F7", "FBEEE4"
SANS, MONO = "IBM Plex Sans", "IBM Plex Mono"


def rgb(h):
    return RGBColor.from_string(h)


# ── helpers ──────────────────────────────────────────────────────────────────
def set_text(shape, new):
    """Replace a shape's text, keeping the first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = new


def find(slide, text):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == text:
            return sh
    raise KeyError(text)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def tb(s, x, y, w, h, text, pt, color=INK, bold=False, font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       italic=False, space=0):
    """Text box; **bold** segments; \n starts a paragraph."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space and i:
            p.space_before = Pt(space)
        for j, seg in enumerate(re.split(r"\*\*", line)):
            if not seg:
                continue
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name, f.size, f.color.rgb = font, Pt(pt), rgb(color)
            f.bold = bold or j % 2 == 1
            f.italic = italic
    return box


def box(s, x, y, w, h, fill=None, line=None, lw=1.0):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def dot(s, x, y, d, color):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2), Inches(y - d / 2), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def arrow(s, x1, y1, x2, y2, color=MUTED, w=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def new_slide(label, title, section, note=None):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb("FFFFFF")
    tb(s, 1.15, 0.92, 19.48 - 1.15, 0.36, label, 18, BLUE, font=MONO)
    tb(s, 1.15, 1.39, 18.33, 0.73, title, 45, INK, bold=True)
    tb(s, 1.15, 10.26, 4.46, 0.36, "DMDU Society Annual Meeting 2026", 18, MUTED)
    tb(s, 15.0, 10.26, 3.94, 0.36, section, 18, MUTED, align=PP_ALIGN.RIGHT)
    if note:
        tb(s, 1.15, 9.81, 18.24, 0.36, note, 18, MUTED)
    return s


def rule(s, x, y, w, color=INK, h=0.02):
    return box(s, x, y, w, h, fill=color)


# ── edits to existing slides ─────────────────────────────────────────────────
# 3 · problem: the fifth box names what the talk varies
set_text(find(S[2], "Architecture + data"), "What the agent sees")

# 14 · becomes the treatment slide, in the Experiment section (slide 20 goes to the appendix)
set_text(find(S[13], "METHOD · THE TOOLS"), "EXPERIMENT · THE TREATMENT")
set_text(find(S[13], "Both arms walk the same network; only v1 sees numbers"),
         "The treatment: v0 sees relations, v1 also sees values")
foot = find(S[13], "Method")
set_text(foot, "Experiment")
foot.left, foot.width = Inches(16.5), Inches(2.44)
foot.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

# 18 · say what the answer key and the hypothesis questions are
set_text(find(S[17], "+ five hypothesis questions (HA1–HE1) · corner number = RDM component"),
         "+ five hypothesis questions (HA1–HE1), each testing one prediction about v0 vs v1 · corner number = RDM "
         "component · answer key example: “half of rural wastewater to septic tanks” = lever level 0.716, "
         "not 0.5")

# 24 · say which comparison the 0/4 → 1/4 is
set_text(find(S[23], "before"), "pilot, before")
set_text(find(S[23], "after rewording the two descriptions"), "pilot, after rewording the two descriptions")
set_text(find(S[23], "Pilot runs, one conversation per arm and question · the full run measures both"),
         "Pilot runs, one conversation per arm and question · full run, v0 vs v1 after rewording: 0 of 10 vs 4 of 10")
set_text(find(S[22], "PILOT FINDINGS"), "DESIGN LESSON 1 · PILOT")
set_text(find(S[23], "PILOT FINDINGS"), "DESIGN LESSON 2 · PILOT")

# 25 · RQ2: pooled numbers, two rows out, a simple bar per arm
s25 = S[24]
set_text(find(s25, "Share of answers, or calls"), "Pooled over all questions")
set_text(find(s25, "Policy level off the key"), "Policy level off the answer key")
set_text(find(s25, "50–87%"), "79%")
set_text(find(s25, "0–7%"), "1%")
set_text(find(s25, "model_network calls per conversation (0–4)"), "model_network calls per conversation")
set_text(find(s25, "2.0–3.1"), "2.3")
set_text(find(s25, "1.1–1.9"), "1.3")
set_text(find(s25, "Level translated literally (designed)"), "Words taken literally (“40%” → 0.4)")
set_text(find(s25, "63%"), "50%")
for sh in list(s25.shapes):               # the old range chart, the last two rows and their rules
    x, y = sh.left / 914400, sh.top / 914400
    if (x >= 11.2 and 2.4 < y < 7.9) or (6.0 <= y < 7.9):
        remove(sh)
set_text(find(s25, "v1 read the levels right and walked the network less. It also did its own arithmetic on the "
                   "values, so more numbers came from no tool."),
         "With values, the agent turned the user’s words into the right setting, and it searched the model less.")
set_text(find(s25, "290 conversations · 5 repetitions per arm · ranges across question groups · first three rows p "
                   "< 0.01"),
         "290 conversations · 5 repetitions per arm · rows 1–3: p < 0.01 in every group tested · uncertainty "
         "ranges: no failures in either arm")
tb(s25, 11.27, 2.53, 7.58, 0.36, "v0 vs v1", 18, MUTED)
rule(s25, 11.27, 2.98, 7.58)
bars = [(0.79, 0.013, "%"), (2.35 / 4, 1.28 / 4, "calls (0–4)"), (0.50, 0.0, "%"), (0.0, 0.40, "%")]
for i, (a, b, _) in enumerate(bars):
    y0 = 3.19 + i * 0.735
    box(s25, 11.27, y0 + 0.03, 7.58, 0.62, fill=PALE)
    for k, (v, col) in enumerate([(a, GREY), (b, BLUE)]):
        yy = y0 + 0.12 + k * 0.24
        if v > 0.005:
            box(s25, 11.27, yy, 7.58 * v, 0.17, fill=col)
        else:
            dot(s25, 11.34, yy + 0.085, 0.17, col)
tb(s25, 11.27, 6.15, 7.58, 0.36, "Bars: share of answers (calls: 0–4 per conversation)", 15, MUTED)

# ── new slides ───────────────────────────────────────────────────────────────
# A · Half is not 0.5
sA = new_slide("RESULTS · RQ2 · ONE EXAMPLE", "“Half” is not 0.5 on the lever", "Findings",
               "Lever levels set in the first run of each conversation · answer keys computed from the model runs, "
               "fixed before the run")
ax0, axw = 6.4, 12.2


def lever_axis(s, y, quote, qid, key, key_note, v0, v0_note, v1, v1_note):
    tb(s, 1.15, y - 0.15, 4.9, 1.6, f"**{qid}**\n“{quote}”", 19.5, INK, space=4)
    rule(s, ax0, y + 0.75, axw, color=GREY, h=0.02)
    for t in (0, 0.5, 1):
        tb(s, ax0 + axw * t - 0.4, y + 0.85, 0.8, 0.3, f"{t:g}", 15, MUTED, align=PP_ALIGN.CENTER)
    kx = ax0 + axw * key
    box(s, kx - 0.015, y + 0.05, 0.03, 0.95, fill=INK)
    tb(s, kx - 6.12, y - 0.32, 6.0, 0.34, f"answer key {key}: {key_note}", 15, INK, bold=True, align=PP_ALIGN.RIGHT)
    for v, n in v0:
        for j in range(n):
            dot(s, ax0 + axw * v, y + 0.56, 0.26, GREY)
    tb(s, ax0, y + 1.2, axw / 2 - 0.2, 0.62, v0_note, 16.5, MUTED2)
    dot(s, ax0 + axw * v1, y + 0.26, 0.26, BLUE)
    tb(s, ax0 + axw / 2, y + 1.2, axw / 2, 0.62, v1_note, 16.5, BLUE, bold=True)


lever_axis(sA, 2.95, "What if half of rural households’ wastewater went to septic tanks by 2050?", "WW1",
           0.716, "half of rural wastewater to septic", [(0.5, 1)], "v0 · 0.5 in 5 of 5: “half” read as "
           "half of the lever", 0.72, "v1 · 0.72 in 5 of 5")
lever_axis(sA, 6.15, "What if Uganda’s landfills captured 40% of the methane their rubbish gives off by 2050?",
           "P50i", 0.688, "40% of landfill gas captured", [(0.4, 1), (0.7, 1), (1.0, 1)],
           "v0 · 0.4 once; 4 times it ran 0.4, 0.7 and 1.0 without choosing", 0.69, "v1 · 0.69 in 5 of 5")
tb(sA, 1.15, 8.75, 18.0, 0.8, "Without values the agent can only guess how words map to settings. With values it "
   "looks them up. Over all 75 policy questions per arm: off the key 79% (v0) vs 1% (v1).", 21, INK)

# B · values invite the agent's own arithmetic
sB = new_slide("RESULTS · RQ2", "Values also invite the agent’s own arithmetic", "Findings",
               "Check 15 re-scored after the run (rounding, lever lists, example thresholds were false positives; "
               "logged) · before the fix: hypothesis 0% vs 32%")
cw = 5.3
steps = [("model_network gave (v1)", "electric light-duty travel\n**20.1 bn km**, at an electric share of **0.61**",
          LIGHT, BLUE),
         ("the agent computed", "“That implies roughly **33 billion km** of light-duty travel in total.”",
          PEACH, ORANGE),
         ("and answered", "“21 ÷ 33 is about **0.64** electric.”\nReasonable, but no tool computed it.",
          PEACH, ORANGE)]
for i, (h, body, fill, col) in enumerate(steps):
    x = 1.15 + i * (cw + 0.55)
    box(sB, x, 3.35, cw, 2.3, fill=fill)
    tb(sB, x + 0.3, 3.55, cw - 0.6, 0.4, h, 16.5, col, bold=True, font=MONO)
    tb(sB, x + 0.3, 4.1, cw - 0.6, 1.5, body, 19.5, INK, space=6)
    if i < 2:
        arrow(sB, x + cw + 0.07, 4.5, x + cw + 0.48, 4.5)
tb(sB, 1.15, 2.6, 17.5, 0.5, "**HA1** · “How far would Uganda need to push the shift of cars and vans to "
   "electricity for them to drive 21 billion kilometres on electricity in 2050?”", 18, MUTED2)
rows = [("Conversations with a number no tool gave", "v0", "v1"),
        ("All 290", "6 of 145", "12 of 145"),
        ("Hypothesis questions (HA1–HE1)", "0%", "20%")]
for r, (a, b, c) in enumerate(rows):
    y = 6.15 + r * 0.6
    head = r == 0
    tb(sB, 1.15, y, 7.5, 0.45, a, 18 if head else 19.5, MUTED if head else INK, bold=head)
    tb(sB, 8.7, y, 1.8, 0.45, b, 18 if head else 19.5, MUTED if head else MUTED2, bold=True)
    tb(sB, 10.6, y, 1.8, 0.45, c, 18 if head else 19.5, MUTED if head else ORANGE, bold=True)
    rule(sB, 1.15, y + 0.5, 11.3, color=INK if head else RULE, h=0.02 if head else 0.01)
tb(sB, 13.0, 6.15, 5.9, 2.4, "Values improve the setting the agent chooses, and give it numbers to calculate with. "
   "Those numbers need tracing before they reach a decision maker.", 21, INK)

# C · RQ1 rewritten: one conversation through the three components
sC = new_slide("RESULTS · RQ1", "Through the agent, a user ran all three RDM steps", "Findings",
               "Q10b → D2b → N1, v1, repetition 1 · all 10 conversations (v0 and v1) took this path · reader "
               "checks pending the hand check")
cols = [("1", "“Is a policy with a strong push on renewable electricity and a moderate cut in deforestation "
               "robust?”", "exploratory_modeling", "500 futures · asks the user what counts as failure"),
        ("2", "“Count it as a failure if 2050 emissions are not at least 5% below BAU.”",
         "vulnerability_analysis\nscenario_discovery", "**fails in 79%** · mostly when GDP grows fast"),
        ("2", "“Does it still work if the economy grows fast?”", "exploratory_modeling\n(GDP top third)\n"
         "vulnerability_analysis", "**fails in 99%**"),
        ("3", "“What could we change so it holds up, and does it?”", "model_network × 3\n"
         "exploratory_modeling\nvulnerability_analysis", "adds a livestock lever: **fails in 0%**")]
cw = 4.15
for i, (comp, q, tools, res) in enumerate(cols):
    x = 1.15 + i * (cw + 0.42)
    tb(sC, x, 2.55, cw, 0.36, f"RDM component {comp}", 16.5, BLUE, font=MONO)
    rule(sC, x, 2.98, cw)
    tb(sC, x, 3.15, cw, 1.45, q, 17, INK, italic=True)
    box(sC, x, 4.7, cw, 1.25, fill=PALE)
    tb(sC, x + 0.2, 4.8, cw - 0.4, 1.1, tools, 15, INK, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    box(sC, x, 6.1, cw, 1.05, fill=LIGHT)
    tb(sC, x + 0.2, 6.15, cw - 0.4, 0.95, res, 18, BLUE, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3:
        arrow(sC, x + cw + 0.04, 5.32, x + cw + 0.38, 5.32)
tb(sC, 1.15, 7.5, 18.0, 0.4, "Across all 290 conversations", 16.5, MUTED, bold=True)
rule(sC, 1.15, 7.92, 17.75, color=RULE, h=0.01)
stats = [("290 / 290", "finished"), ("0", "loops"), ("0", "invented thresholds"), ("1", "forecast language"),
         ("3", "named a winner"), ("22% → 12%", "required RDM step skipped, v0 → v1")]
sw = 17.75 / len(stats)
for i, (n, lab) in enumerate(stats):
    x = 1.15 + i * sw
    tb(sC, x, 8.05, sw - 0.2, 0.55, n, 27, BLUE if i in (0, 5) else INK, bold=True)
    tb(sC, x, 8.65, sw - 0.2, 0.7, lab, 15, MUTED)

# D · conclusion
sD = new_slide("CONCLUSION", "Design what the agent sees, and test it", "Findings")
lessons = [("Show values, not only links", "With the model’s values the agent turns words into the right "
            "setting (79% → 1% off the key) and searches less."),
           ("Trace and label every number", "Values invite the agent’s own arithmetic, and one unlabeled number "
            "misled it in half the conversations. Label each source; trace what the agent computes."),
           ("Tool wording chooses the RDM step", "One sentence in a tool description decided whether the agent "
            "compared policies with trade_off_analysis. Test it like a model.")]
cw = 5.6
for i, (h, body) in enumerate(lessons):
    x = 1.15 + i * (cw + 0.47)
    box(sD, x, 2.6, cw, 4.0, fill=BLUE if i == 0 else LIGHT)
    tb(sD, x + 0.35, 2.85, cw - 0.7, 0.5, str(i + 1), 30, "FFFFFF" if i == 0 else BLUE, bold=True)
    tb(sD, x + 0.35, 3.55, cw - 0.7, 1.0, h, 25.5, "FFFFFF" if i == 0 else INK, bold=True)
    tb(sD, x + 0.35, 4.75, cw - 0.7, 1.8, body, 18, "FFFFFF" if i == 0 else INK)
tb(sD, 1.15, 7.05, 8.6, 0.4, "What we have not shown", 16.5, MUTED, bold=True)
tb(sD, 1.15, 7.5, 8.6, 1.8, "One model and one country · one agent model · scripted users, not policy "
   "makers · 5 repetitions per arm · the reader model is checked on 20 judgments", 18, MUTED2)
tb(sD, 10.3, 7.05, 8.6, 0.4, "Next", 16.5, MUTED, bold=True)
tb(sD, 10.3, 7.5, 8.6, 1.8, "Sessions with real users in Uganda · checking the agent’s explanations of why "
   "a policy works · tracing the agent’s own numbers in the answer", 18, MUTED2)

# E · appendix divider
sE = new_slide("APPENDIX", "Appendix: details for questions", "Appendix")

# ── order: drop slide 7, move four method slides and slide 20 to the appendix ─
lst = prs.slides._sldIdLst
ids = list(lst)
new = {k: ids[len(S) + i] for i, k in enumerate("ABCDE")}
order = [0, 1, 2, 3, 4,                       # context
         5, 7, 9, 11, 15,                     # method: surrogate, harness, tools, network
         16, 13, 17, 18, 20,                  # experiment: treatment, questions, paths, run design
         21, "A", 24, "B", "C", 22, 23, "D",  # findings
         "E", 8, 10, 12, 14, 19]              # appendix
dropped = [6, 25]                             # surrogate placeholder (slide 8 says it) · old RQ1 (replaced by C)
for el in ids:
    lst.remove(el)
for k in order:
    lst.append(ids[k] if isinstance(k, int) else new[k])
for k in dropped:
    prs.part.drop_rel(ids[k].get(qn("r:id")))
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
