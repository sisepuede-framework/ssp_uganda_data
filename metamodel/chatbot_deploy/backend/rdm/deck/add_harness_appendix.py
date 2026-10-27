# Requires python-pptx. Usage: python backend/rdm/deck/add_harness_appendix.py IN.pptx OUT.pptx
"""Append the harness text, word for word, to a deck exported from Claude Design (IBM Plex, 20 × 11.25 in).

Not part of the talk: everything the agent is told, for review of the instructions. Prompt, tool descriptions,
loop settings and rule messages come from harness_text.json (refresh with
python -m backend.rdm.deck.dump_harness_text); the fixed notes inside the replies are copied by hand from
backend/rdm/tools/*.py (checked against the code 2026-10-27).
"""
import json
import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HARNESS = json.loads((Path(__file__).parent / "harness_text.json").read_text())
BLUE, INK, MUTED = "1F4E9C", "16181D", "6B707B"
SANS, MONO = "IBM Plex Sans", "IBM Plex Mono"
LM, CW, TOP, BOTTOM = 1.15, 17.79, 2.45, 9.65
LABEL = "APPENDIX · THE HARNESS, WORD FOR WORD"
NOTE = ("Word for word what the agent receives, from the code (backend/rdm/deck/harness_text.json, written by "
        "dump_harness_text.py). Not part of the talk.")


def tb(s, x, y, w, h, text, pt, color=INK, bold=False, font=SANS, align=PP_ALIGN.LEFT, space=0):
    """Text box; **bold** segments; \n starts a paragraph."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space and i:
            p.space_before = Pt(space)
        for j, seg in enumerate(re.split(r"\*\*", line)):
            if seg:
                r = p.add_run()
                r.text = seg
                r.font.name, r.font.size, r.font.color.rgb = font, Pt(pt), RGBColor.from_string(color)
                r.font.bold = bold or j % 2 == 1
    return box


def new_slide(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    tb(s, LM, 0.92, CW, 0.36, LABEL, 18, BLUE, font=MONO)
    tb(s, LM, 1.39, 18.33, 0.73, title, 36, INK, bold=True)
    tb(s, LM, 10.26, 4.46, 0.36, "DMDU Society Annual Meeting 2026", 18, MUTED)
    tb(s, 15.0, 10.26, 3.94, 0.36, "Appendix", 18, MUTED, align=PP_ALIGN.RIGHT)
    tb(s, LM, 9.81, CW, 0.36, NOTE, 15, MUTED)
    return s


def flow(prs, title, paras, ncols=2, pt=15):
    """Pour paragraphs into columns, continuing on new slides when they fill up."""
    gap = 0.5
    col_w = (CW - (ncols - 1) * gap) / ncols
    cpl = col_w * 72 / (pt * 0.44)                 # characters per line, IBM Plex Sans
    cap = 0.93 * (BOTTOM - TOP) * 72 / (pt * 1.24) # lines per column, with a margin
    cols, cur, used = [], [], 0.0
    for p in paras:
        n = max(1, math.ceil(len(p.replace("**", "")) / cpl)) + 0.3
        if cur and used + n > cap:
            cols.append(cur)
            cur, used = [], 0.0
        cur.append(p)
        used += n
    cols.append(cur)
    for k in range(0, len(cols), ncols):
        s = new_slide(prs, title + (" (cont.)" if k else ""))
        for c, col in enumerate(cols[k:k + ncols]):
            tb(s, LM + c * (col_w + gap), TOP, col_w, BOTTOM - TOP, "\n".join(col), pt, space=4)


def prompt_paras(block):
    out = []
    for ln in block.split("\n"):
        if ln.strip():
            out.append(f"**{ln[3:]}**" if ln.startswith("## ") else ("• " + ln[2:] if ln.startswith("- ") else ln))
    return out


def param_line(name, spec, required):
    extra = f"; one of: {', '.join(map(str, spec['enum']))}" if "enum" in spec else ""
    req = ", required" if name in required else ""
    desc = spec.get("description", "")
    return f"   • **{name}** ({spec.get('type', '')}{req}{extra}){': ' + desc if desc else ''}"


def append(prs):
    s = new_slide(prs, "Appendix: the text of the harness")
    tb(s, LM, TOP, CW, 0.5, "Not part of the presentation. Everything the agent is told, word for word, for review "
       "of the instructions. Identical in v0 and v1, except what model_network returns.", 21, BLUE, bold=True)
    parts = [("System prompt", "method, how the model is organized, the tools, the rules, the catalog",
              "harness/prompt.py"),
             ("Tool descriptions", "the 6 tools and their parameters, as the agent sees them", "harness/registry.py"),
             ("The loop and the scripted user", "model, limits, when the scripted user replies and with what",
              "harness/agent.py · eval/questions.py"),
             ("Rules checked in code", "the messages the tools return when a call breaks a rule",
              "tools/technical.py · information.py · common.py"),
             ("Fixed text inside the replies", "notes and requests the tools add to their results", "tools/*.py")]
    for i, (head, sub, path) in enumerate(parts):
        y = TOP + 1.3 + i * 1.15
        tb(s, LM, y, 5.6, 0.5, head, 21, INK, bold=True)
        tb(s, LM + 5.9, y + 0.04, 7.6, 0.5, sub, 18, INK)
        tb(s, LM + 13.8, y + 0.06, CW - 13.8, 0.5, path, 15, MUTED, font=MONO)

    body, cat = HARNESS["system_prompt"].split("## Catalogue")
    flow(prs, "System prompt", prompt_paras(body), ncols=2, pt=14)
    flow(prs, "System prompt: the catalogue (end of the prompt)", prompt_paras(cat), ncols=4, pt=10)

    paras = []
    for t in HARNESS["tools"]:
        sch = t["input_schema"]
        paras += [f"**{t['name']}**", t["description"]]
        paras += [param_line(n, sp, sch.get("required", [])) for n, sp in sch["properties"].items()]
    flow(prs, "Tool descriptions (the 6 tools, every arm)", paras, ncols=3, pt=13.5)

    lp = HARNESS["loop"]
    flow(prs, "The loop and the scripted user", [
        "**One conversation = one question, in one arm, one repetition.** The agent chooses the tools; the harness "
        "runs them and never chooses for it.",
        f"**Model:** {lp['model']}, effort {lp['effort']}, the same in every arm. At most {lp['max_tokens']:,} output "
        f"tokens per request and {lp['max_steps']} requests per conversation. Reasoning is shown as a summary "
        "(display only; same cost).",
        "**Arms:** v0 (relations only), v1 (relations + values). Tool names, schemas and descriptions are identical; "
        "only what model_network returns differs.",
        "**The scripted user** sends the question, then the question's scripted replies in order whenever the agent "
        "ends its turn. When none are left and the agent's message ends with \"?\", it replies once with the neutral "
        "go-ahead; otherwise the conversation ends.",
        f"**Neutral go-ahead:** \"{lp['go_ahead']}\"",
        "**No model fallback:** a refusal is recorded as such, because switching model mid-run would confound the "
        "arms.",
        "**exploratory_modeling** always draws its futures with the same seed (0), so repeated calls see the same "
        "futures.",
        "**Logged per tool call:** turn, tool, arguments, status (ok / error), seconds, arm, reply size and the full "
        "reply; per request: stop reason and tokens. The final answer is the agent's last text.",
    ], ncols=1, pt=19.5)

    errs = [e for e in HARNESS["tool_errors"] if not e["message"].startswith('f"arm must')]
    flow(prs, "Rules checked in code: what the agent reads when a call breaks one", [
        "When a call breaks a rule the tool returns {\"error\": message} and the agent can try again. Messages as "
        "written in the code ({…} is filled in at run time):"] +
        [f"• {e['message'].lstrip('f').strip(chr(34))}   ({e['file']}:{e['line']})" for e in errs] +
        ["• No ensemble {ensemble_id!r}; known: {…}. Run exploratory_modeling first.   (harness/state.py)",
         "• Unknown tool {name!r}; tools: [the 6 names]   (harness/registry.py)"], ncols=2, pt=15)

    flow(prs, "Fixed text inside the tool replies", [
        "**vulnerability_analysis, no threshold given** → status needs_threshold, with: \"Which value of {metric} in "
        "{year} counts as failure for {policy!r}? Give it in {unit}.\" plus the distribution to help choose.",
        "**exploratory_modeling, scenario_discovery** → \"Uncertainties are positions in their sampled range: 0 = "
        "low end, 0.5 = central, 1 = high end.\"",
        "**scenario_discovery** → \"Balanced accuracy of a small tree separating pass from fail (0.5 = no better "
        "than chance). The gap between the two is what the other uncertainties add.\"",
        "**trade_off_analysis** → \"Changes are against BAU in the same futures. Cost is the sum of the technical, "
        "system and fuel cost types (money spent, positive = more spending); net benefit adds every benefit and "
        "subtracts the cost (positive = better for the country). Cost per tonne counts only the cost types …\" + "
        "the costs note below",
        "**exploratory_modeling (cost or benefit metrics), trade_off_analysis** → \"Costs and benefits are shown only "
        "as changes against BAU on the same futures: their levels are differences against one fixed reference run, "
        "not a policy's money. Every cost_* is money spent; benefit_* and net_benefit are good for the country. "
        "cost_pct_gdp changes are in percentage points of GDP. cost_fuel holds only biogas and waste-to-energy "
        "revenue; fuel spending is in benefit_sector_specific … Costs and benefits that do not depend on a lever's "
        "level cancel out and do not show.\"",
        "**model_network, every arm** → a category: \"the sum of its emission fields; exploratory_modeling reports "
        "this metric\" · NemoMod: \"SISEPUEDE's electricity optimisation, one box: what goes in and what comes out "
        "are known, not the steps inside\" · a lever with no chain: \"Its inputs are not used to compute anything in "
        "the traced run of the model code (a gap of the network), so no chain is known for it.\" · third given with "
        "a lever only: \"third filters an uncertainty; with a lever alone, the node is shown for each of the lever's "
        "thirds (low, mid, high)\" · long lists: \"{n} more fields; ask for one by its field name\"",
        "**model_network, v1 only** → what a lever sets: \"what the lever level means: the input at BAU (bottom of "
        "the range run), the middle and the top of the range run, a straight line in between\" · values: \"Median "
        "and p10–p90 over the experiment's runs split by L{id}'s level: low 0–⅓, mid ⅓–⅔, high ⅔–1 of its range[, "
        "only futures with X{id} ({name}) in its {third} third]; every other lever and uncertainty varies. Values "
        "from the model runs, not the surrogate.\" · \"L{id} does not reach this variable: values over every lever "
        "level\" · from a category: each lever that reaches it, with the median at its low and high third and the "
        "change.",
        "**country_context** → \"Modelled baseline assumptions (2025–2050), not observed present-day statistics.\"  "
        "·  GDP topic, added to the reply: \"These values are SISEPUEDE's base future (future 0). It is not one of "
        "the experiment's runs, and its GDP is above the 90th percentile of the runs in 2040 and 2050. Do not read "
        "it as the middle of the futures explored.\"",
    ], ncols=2, pt=15)


if __name__ == "__main__":
    prs = Presentation(sys.argv[1])
    n = len(prs.slides)
    append(prs)
    prs.save(sys.argv[2])
    print("saved", sys.argv[2], f"{n} + {len(prs.slides) - n} appendix slides")
